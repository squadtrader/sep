#!/usr/bin/env python3
"""
boc_mpr_scraper.py
-------------------
Collecte les projections économiques publiées dans le Rapport sur la politique
monétaire (RPM / MPR) de la Banque du Canada, à chaque publication trimestrielle
(janvier, avril, juillet, octobre).

Contrairement au SEP de la Fed (médiane / central tendency / range par participant),
la Banque du Canada publie un scénario de référence unique ("base-case projection"),
avec les chiffres du rapport précédent entre parenthèses pour comparaison. Deux
tableaux sont extraits de la page "Projections" de chaque RPM :

  - Tableau 2 : Contributions à la croissance annuelle du PIB réel (Consommation,
    Logement, Gouvernement, Investissement des entreprises, Exportations,
    Importations, Stocks, PIB total, output potentiel, inflation IPC) — un chiffre
    par année.
  - Tableau 3 : Résumé de la projection trimestrielle pour le Canada (inflation IPC,
    inflation core, PIB réel en glissement annuel et trimestriel) — détail
    trimestriel à court terme + chiffres annuels (Q4/Q4) à plus long terme.

Étapes :
1. Parcourt les pages de la liste des RPM (bankofcanada.ca/publications/mpr/) pour
   trouver les URL de chaque rapport (format "mpr-AAAA-MM-JJ").
2. Pour chaque rapport, télécharge sa page "/projections/" et en extrait les
   Tableaux 2 et 3.
3. Sauvegarde le tout en JSON, avec exports CSV et Excel optionnels.

Usage :
    python boc_mpr_scraper.py                          # tous les RPM trouvés
    python boc_mpr_scraper.py --since 2024              # seulement à partir de 2024
    python boc_mpr_scraper.py --latest                  # seulement le RPM le plus récent
    python boc_mpr_scraper.py --out boc_mpr.json --xlsx boc_mpr.xlsx
    python boc_mpr_scraper.py --cache-dir .cache_boc     # évite de re-télécharger

Note : la Banque du Canada a renouvelé son site en 2024-2025. Les RPM antérieurs à
2024 n'utilisent pas toujours le même schéma d'URL / la même page "Projections" ;
le script les ignore proprement (avec un message) plutôt que de planter dessus.

Dépendances : requests, beautifulsoup4, openpyxl
    pip install requests beautifulsoup4 openpyxl
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

LISTING_URL_TMPL = "https://www.bankofcanada.ca/publications/mpr/?mt_page={page}"
MAX_LISTING_PAGES = 15  # filet de sécurité pour ne pas boucler indéfiniment

# La BOC ne publie pas de projection de son propre taux directeur (contrairement
# au "dot plot" de la Fed) : ses projections de croissance/inflation supposent que
# le taux suit le chemin anticipé par les marchés, mais ce chemin n'est pas publié
# dans les tableaux du RPM. On récupère à la place, via l'API Valet, le taux
# directeur RÉELLEMENT en vigueur à la date de chaque RPM (donnée factuelle, pas
# une projection) pour donner un repère.
VALET_OVERNIGHT_RATE_URL_TMPL = (
    "https://www.bankofcanada.ca/valet/observations/V39079/json"
    "?start_date={start}&end_date={end}"
)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    )
}

# Format d'URL utilisé depuis ~2024 : /publications/mpr/mpr-AAAA-MM-JJ/
REPORT_URL_RE = re.compile(
    r"https://www\.bankofcanada\.ca/publications/mpr/(mpr-\d{4}-\d{2}-\d{2})/?\""
)

TABLE2_ROW_PATTERNS = [
    ("consumption", re.compile(r"^consumption$", re.I)),
    ("housing", re.compile(r"^housing$", re.I)),
    ("government", re.compile(r"^government$", re.I)),
    ("business_investment", re.compile(r"business fixed investment", re.I)),
    ("final_domestic_demand", re.compile(r"final domestic demand", re.I)),
    ("exports", re.compile(r"^exports$", re.I)),
    ("imports", re.compile(r"^imports$", re.I)),
    ("inventories", re.compile(r"^inventories$", re.I)),
    ("gdp", re.compile(r"^gdp$", re.I)),
    ("potential_output_range", re.compile(r"potential output", re.I)),
    ("cpi_inflation_annual", re.compile(r"cpi inflation", re.I)),
]

TABLE3_ROW_PATTERNS = [
    ("cpi_inflation", re.compile(r"^cpi inflation", re.I)),
    ("core_inflation", re.compile(r"^core inflation", re.I)),
    ("real_gdp_yoy", re.compile(r"real gdp \(year", re.I)),
    ("real_gdp_qoq", re.compile(r"real gdp \(quarter", re.I)),
]


@dataclass
class MprReport:
    date: str  # AAAA-MM-JJ
    url: str
    table2: dict = field(default_factory=dict)   # variable -> year -> {"value":.., "prior":..}
    table3_quarterly: dict = field(default_factory=dict)  # variable -> "AAAA-QN" -> {...}
    table3_annual: dict = field(default_factory=dict)     # variable -> year -> {...}
    policy_rate: Optional[str] = None  # taux directeur réel (%) en vigueur à la date du RPM


def fetch(url: str, cache_dir: Optional[Path] = None) -> str:
    if cache_dir:
        cache_dir.mkdir(parents=True, exist_ok=True)
        cache_file = cache_dir / (re.sub(r"[^a-zA-Z0-9]+", "_", url) + ".html")
        if cache_file.exists():
            return cache_file.read_text(encoding="utf-8")

    resp = requests.get(url, headers=HEADERS, timeout=30)
    resp.raise_for_status()
    resp.encoding = "utf-8"  # le site ne déclare pas toujours son charset -> forcer UTF-8
    html = resp.text

    if cache_dir:
        cache_file.write_text(html, encoding="utf-8")

    return html


def discover_report_dates(cache_dir: Optional[Path] = None, max_pages: int = MAX_LISTING_PAGES) -> list[str]:
    """Parcourt les pages de la liste des RPM et retourne les dates (AAAA-MM-JJ)
    de tous les rapports utilisant le schéma d'URL moderne."""
    dates: set[str] = set()
    for page in range(1, max_pages + 1):
        url = LISTING_URL_TMPL.format(page=page)
        try:
            html = fetch(url, cache_dir=cache_dir)
        except requests.HTTPError:
            break

        found = set(m.group(1)[4:] for m in REPORT_URL_RE.finditer(html))
        if not found or found <= dates:
            # Rien de nouveau sur cette page -> on a atteint la fin de la liste utile
            break
        dates.update(found)
        time.sleep(0.5)

    return sorted(dates)


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _split_value_prior(text: str) -> tuple[Optional[str], Optional[str]]:
    """Sépare une cellule "2.2 (2.2)" en (valeur_actuelle, valeur_rapport_précédent).
    Gère aussi les valeurs simples "2.5" (pas de comparaison) et les cellules vides."""
    text = _clean(text)
    if not text or text in {"-", "—", ".."}:
        return None, None
    m = re.match(r"^(.*?)(?:\s*\(([^()]+)\))?$", text)
    if not m:
        return text, None
    value = m.group(1).strip()
    prior = m.group(2).strip() if m.group(2) else None
    return (value or None), prior


def _expand_row(row) -> list[str]:
    """Étend une ligne d'en-tête en tenant compte des colspan (une entrée par colonne)."""
    out = []
    for cell in row.find_all(["th", "td"]):
        text = _clean(cell.get_text(" "))
        colspan = int(cell.get("colspan", 1))
        out.extend([text] * colspan)
    return out


def _find_table_by_content(soup: BeautifulSoup, patterns: list, min_matches: int):
    """Repère la <table> voulue par son contenu (libellés de lignes attendus) plutôt
    que par son numéro ("Table 2", "Table 3"), qui varie d'une édition du RPM à
    l'autre selon le nombre de graphiques qui la précèdent sur la page. On compare
    les patterns au libellé de chaque ligne (première cellule), pas au texte
    concaténé de toute la table, car les patterns sont ancrés (^...$)."""
    best_table, best_score = None, 0
    for table in soup.find_all("table"):
        matched_keys = set()
        for row in table.find_all("tr"):
            cells = row.find_all(["th", "td"])
            if not cells:
                continue
            label = _clean(cells[0].get_text(" "))
            if not label:
                continue
            for key, pattern in patterns:
                if pattern.search(label):
                    matched_keys.add(key)
        if len(matched_keys) > best_score:
            best_table, best_score = table, len(matched_keys)
    return best_table if best_score >= min_matches else None


def parse_table2(html: str) -> dict:
    """Tableau des contributions à la croissance annuelle du PIB réel — repéré par
    son contenu (Consumption/Housing/Government/GDP...), pas par son numéro."""
    soup = BeautifulSoup(html, "html.parser")
    table = _find_table_by_content(soup, TABLE2_ROW_PATTERNS, min_matches=4)
    if table is None:
        raise ValueError("Tableau des contributions au PIB introuvable.")

    rows = table.find_all("tr")
    header_cells = _expand_row(rows[0])
    years = header_cells[1:]  # on saute la cellule de libellé

    result: dict = {}
    for row in rows[1:]:
        cells = row.find_all(["th", "td"])
        if not cells:
            continue
        texts = [_clean(c.get_text(" ")) for c in cells]
        if not texts or not texts[0]:
            continue
        label = texts[0]

        matched_key = None
        for key, pattern in TABLE2_ROW_PATTERNS:
            if pattern.search(label):
                matched_key = key
                break
        if not matched_key:
            continue

        entry: dict = {}
        for year, cell_text in zip(years, texts[1:]):
            value, prior = _split_value_prior(cell_text)
            if value is not None:
                entry[year] = {"value": value, "prior": prior}
        result[matched_key] = {"label": label, "years": entry}

    return result


def parse_table3(html: str) -> tuple[dict, dict]:
    """Tableau 3 : Résumé de la projection trimestrielle — deux lignes d'en-tête
    (année puis trimestre), avec une colonne "vide" séparant le détail trimestriel
    à court terme des chiffres annuels (Q4/Q4) à plus long terme."""
    soup = BeautifulSoup(html, "html.parser")
    table = _find_table_by_content(soup, TABLE3_ROW_PATTERNS, min_matches=3)
    if table is None:
        raise ValueError("Tableau de résumé trimestriel introuvable.")

    rows = table.find_all("tr")
    if len(rows) < 3:
        raise ValueError("Tableau 3 : structure d'en-tête inattendue.")

    year_row = _expand_row(rows[0])[1:]   # on saute la cellule de libellé (rowspan)
    quarter_row = _expand_row(rows[1])    # pas de cellule de libellé sur cette ligne

    n = min(len(year_row), len(quarter_row))
    year_row, quarter_row = year_row[:n], quarter_row[:n]

    # Séparer le détail trimestriel (court terme) des chiffres annuels Q4/Q4
    # (plus long terme). Deux indices possibles selon l'édition du RPM :
    #  1) une colonne "vide" (les deux en-têtes vides) sépare visuellement les deux blocs ;
    #  2) le bloc annuel se reconnaît à une suite d'années strictement consécutives
    #     en fin de tableau (ex: 2025, 2026, 2027, 2028), qu'il y ait ou non une
    #     colonne vide entre les deux blocs.
    gap_idx = None
    for i, (y, q) in enumerate(zip(year_row, quarter_row)):
        if not y and not q:
            gap_idx = i
            break

    quarterly_cols: list[int] = []
    annual_cols: list[int] = []

    if gap_idx is not None:
        quarterly_cols = [i for i in range(gap_idx) if year_row[i] and quarter_row[i]]
        annual_cols = [i for i in range(gap_idx + 1, n) if year_row[i]]
    else:
        # Chercher le plus long suffixe d'années strictement consécutives (>= 2 colonnes)
        annual_start = n
        for i in range(n - 1, 0, -1):
            try:
                y_cur, y_prev = int(year_row[i]), int(year_row[i - 1])
            except ValueError:
                break
            if y_cur == y_prev + 1:
                annual_start = i - 1
            else:
                break
        if annual_start < n - 1:  # au moins 2 colonnes dans le suffixe -> bloc annuel trouvé
            annual_cols = [i for i in range(annual_start, n) if year_row[i]]
            quarterly_cols = [i for i in range(annual_start) if year_row[i] and quarter_row[i]]
        else:
            quarterly_cols = [i for i in range(n) if year_row[i] and quarter_row[i]]

    quarterly: dict = {}
    annual: dict = {}

    for row in rows[2:]:
        cells = row.find_all(["th", "td"])
        if not cells:
            continue
        texts = [_clean(c.get_text(" ")) for c in cells]
        if not texts or not texts[0]:
            continue
        label = texts[0]
        data_cells = texts[1:1 + n]

        matched_key = None
        for key, pattern in TABLE3_ROW_PATTERNS:
            if pattern.search(label):
                matched_key = key
                break
        if not matched_key:
            continue

        q_entry: dict = {}
        for i in quarterly_cols:
            if i >= len(data_cells):
                continue
            value, prior = _split_value_prior(data_cells[i])
            if value is not None:
                period = f"{year_row[i]}-{quarter_row[i]}"
                q_entry[period] = {"value": value, "prior": prior}

        a_entry: dict = {}
        for i in annual_cols:
            if i >= len(data_cells):
                continue
            value, prior = _split_value_prior(data_cells[i])
            if value is not None:
                a_entry[year_row[i]] = {"value": value, "prior": prior}

        if q_entry:
            quarterly[matched_key] = {"label": label, "periods": q_entry}
        if a_entry:
            annual[matched_key] = {"label": label, "years": a_entry}

    return quarterly, annual


def fetch_policy_rate(date: str, cache_dir: Optional[Path] = None) -> Optional[str]:
    """Récupère le taux directeur (cible du taux à un jour) réellement en vigueur
    à la date d'un RPM, via l'API Valet (série V39079). On interroge une fenêtre
    de 30 jours avant la date pour être sûr d'obtenir une valeur (la série n'a une
    entrée que les jours ouvrables, et le taux n'est ajusté qu'à date fixe)."""
    end = datetime.strptime(date, "%Y-%m-%d")
    start = end - timedelta(days=30)
    url = VALET_OVERNIGHT_RATE_URL_TMPL.format(start=start.strftime("%Y-%m-%d"), end=date)

    try:
        raw = fetch(url, cache_dir=cache_dir)
        data = json.loads(raw)
        observations = data.get("observations", [])
        if not observations:
            return None
        last = observations[-1]
        return last.get("V39079", {}).get("v")
    except (requests.HTTPError, json.JSONDecodeError, KeyError, ValueError):
        return None


def scrape_report(date: str, cache_dir: Optional[Path] = None, pause: float = 1.0) -> Optional[MprReport]:
    base_url = f"https://www.bankofcanada.ca/publications/mpr/mpr-{date}/"
    proj_url = base_url + "projections/"

    try:
        html = fetch(proj_url, cache_dir=cache_dir)
    except requests.HTTPError as exc:
        print(f"  ! {date}: page 'Projections' introuvable ({exc})", file=sys.stderr)
        return None

    report = MprReport(date=date, url=proj_url)

    try:
        report.table2 = parse_table2(html)
    except ValueError as exc:
        print(f"  ! {date}: Tableau 2 non extrait ({exc})", file=sys.stderr)

    try:
        report.table3_quarterly, report.table3_annual = parse_table3(html)
    except ValueError as exc:
        print(f"  ! {date}: Tableau 3 non extrait ({exc})", file=sys.stderr)

    report.policy_rate = fetch_policy_rate(date, cache_dir=cache_dir)

    if not report.table2 and not report.table3_annual:
        return None

    time.sleep(pause)
    return report


def report_to_dict(r: MprReport) -> dict:
    return {
        "date": r.date,
        "url": r.url,
        "policy_rate": r.policy_rate,
        "table2_contributions": r.table2,
        "table3_quarterly": r.table3_quarterly,
        "table3_annual": r.table3_annual,
    }


def write_json(reports: list[MprReport], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    payload = [report_to_dict(r) for r in reports]
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(reports: list[MprReport], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["report_date", "table", "variable", "label", "period", "value", "prior_report_value"])
        for r in reports:
            for var_key, var_data in r.table2.items():
                for year, vals in var_data.get("years", {}).items():
                    writer.writerow([r.date, "table2", var_key, var_data.get("label", var_key), year,
                                      vals["value"], vals.get("prior") or ""])
            for var_key, var_data in r.table3_annual.items():
                for year, vals in var_data.get("years", {}).items():
                    writer.writerow([r.date, "table3_annual", var_key, var_data.get("label", var_key), year,
                                      vals["value"], vals.get("prior") or ""])
            for var_key, var_data in r.table3_quarterly.items():
                for period, vals in var_data.get("periods", {}).items():
                    writer.writerow([r.date, "table3_quarterly", var_key, var_data.get("label", var_key), period,
                                      vals["value"], vals.get("prior") or ""])


# --- Export Excel -----------------------------------------------------------

_XLSX_FONT_NAME = "Arial"
_XLSX_HEADER_FILL = PatternFill("solid", fgColor="A6192E")  # rouge Banque du Canada
_XLSX_HEADER_FONT = Font(name=_XLSX_FONT_NAME, size=10, bold=True, color="FFFFFF")
_XLSX_SUBHEADER_FILL = PatternFill("solid", fgColor="F2D7D9")
_XLSX_SUBHEADER_FONT = Font(name=_XLSX_FONT_NAME, size=10, bold=True)
_XLSX_CELL_FONT = Font(name=_XLSX_FONT_NAME, size=10)
_XLSX_THIN = Side(style="thin", color="B7B7B7")
_XLSX_BORDER = Border(left=_XLSX_THIN, right=_XLSX_THIN, top=_XLSX_THIN, bottom=_XLSX_THIN)
_XLSX_CENTER = Alignment(horizontal="center", vertical="center")

SUMMARY_VARIABLES = [
    ("cpi_inflation", "Inflation IPC (T4/T4)"),
    ("core_inflation", "Inflation core (T4/T4)"),
    ("real_gdp_yoy", "PIB réel (glissement annuel)"),
]


def _xlsx_style_header(cell, fill, font) -> None:
    cell.fill, cell.font, cell.alignment, cell.border = fill, font, _XLSX_CENTER, _XLSX_BORDER


def _xlsx_autofit(ws, min_width=10, max_width=26) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, min_width), max_width)


def _fmt_cell(vals: Optional[dict]) -> str:
    if not vals:
        return ""
    value, prior = vals.get("value"), vals.get("prior")
    return f"{value} ({prior})" if prior else (value or "")


def write_xlsx(reports: list[MprReport], out_path: Path) -> None:
    if not reports:
        return
    reports = sorted(reports, key=lambda r: r.date)

    years: set = set()
    for r in reports:
        for var_data in r.table3_annual.values():
            years.update(var_data.get("years", {}).keys())
    years_sorted = sorted(years, key=lambda y: int(y))

    wb = Workbook()

    # --- Feuille résumé : inflation IPC / core / PIB réel, par année, valeur (rapport précédent)
    ws = wb.active
    ws.title = "Résumé (annuel Q4-Q4)"
    ws.cell(row=1, column=1, value="Date du RPM")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    _xlsx_style_header(ws.cell(row=1, column=1), _XLSX_HEADER_FILL, _XLSX_HEADER_FONT)

    ws.cell(row=1, column=2, value="Taux directeur (%)*")
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    _xlsx_style_header(ws.cell(row=1, column=2), _XLSX_HEADER_FILL, _XLSX_HEADER_FONT)

    col = 3
    for var_key, var_label in SUMMARY_VARIABLES:
        start_col = col
        for year in years_sorted:
            c = ws.cell(row=2, column=col, value=year)
            _xlsx_style_header(c, _XLSX_SUBHEADER_FILL, _XLSX_SUBHEADER_FONT)
            col += 1
        if col == start_col:
            # Aucune année trouvée pour cette variable (donnée manquante sur le
            # site) -> une seule colonne "N/D" plutôt que de planter sur une
            # fusion de cellules invalide (0 colonne).
            c = ws.cell(row=2, column=col, value="N/D")
            _xlsx_style_header(c, _XLSX_SUBHEADER_FILL, _XLSX_SUBHEADER_FONT)
            col += 1
        if col - 1 > start_col:
            ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=col - 1)
        _xlsx_style_header(ws.cell(row=1, column=start_col, value=var_label), _XLSX_HEADER_FILL, _XLSX_HEADER_FONT)

    row = 3
    for r in reports:
        cell = ws.cell(row=row, column=1, value=r.date)
        cell.font, cell.border, cell.alignment = _XLSX_CELL_FONT, _XLSX_BORDER, _XLSX_CENTER

        rate_cell = ws.cell(row=row, column=2, value=r.policy_rate or "")
        rate_cell.font, rate_cell.border, rate_cell.alignment = _XLSX_CELL_FONT, _XLSX_BORDER, _XLSX_CENTER

        col = 3
        for var_key, _ in SUMMARY_VARIABLES:
            var_data = r.table3_annual.get(var_key, {}).get("years", {})
            for year in years_sorted:
                c = ws.cell(row=row, column=col, value=_fmt_cell(var_data.get(year)))
                c.font, c.border, c.alignment = _XLSX_CELL_FONT, _XLSX_BORDER, _XLSX_CENTER
                col += 1
        row += 1
    ws.freeze_panes = "C3"
    _xlsx_autofit(ws)

    note_row = row + 1
    ws.cell(row=note_row, column=1,
            value="* Taux directeur réellement en vigueur à la date du RPM (source : API Valet, série V39079) "
                  "— la BOC ne publie pas de projection de son propre taux, contrairement à la Fed.")
    ws.cell(row=note_row, column=1).font = Font(name=_XLSX_FONT_NAME, size=9, italic=True)

    # --- Feuille détail : contributions au PIB (Tableau 2)
    ws2 = wb.create_sheet("Contributions PIB (T2)")
    t2_years: set = set()
    for r in reports:
        for var_data in r.table2.values():
            t2_years.update(var_data.get("years", {}).keys())
    t2_years_sorted = sorted(t2_years, key=lambda y: int(y))

    ws2.cell(row=1, column=1, value="Date du RPM")
    ws2.cell(row=1, column=2, value="Composante")
    for j, year in enumerate(t2_years_sorted, start=3):
        _xlsx_style_header(ws2.cell(row=1, column=j, value=year), _XLSX_HEADER_FILL, _XLSX_HEADER_FONT)
    _xlsx_style_header(ws2.cell(row=1, column=1), _XLSX_HEADER_FILL, _XLSX_HEADER_FONT)
    _xlsx_style_header(ws2.cell(row=1, column=2), _XLSX_HEADER_FILL, _XLSX_HEADER_FONT)

    row = 2
    for r in reports:
        for var_key, _ in TABLE2_ROW_PATTERNS:
            var_data = r.table2.get(var_key)
            if not var_data:
                continue
            ws2.cell(row=row, column=1, value=r.date).font = _XLSX_CELL_FONT
            ws2.cell(row=row, column=2, value=var_data.get("label", var_key)).font = _XLSX_CELL_FONT
            for j, year in enumerate(t2_years_sorted, start=3):
                val = _fmt_cell(var_data.get("years", {}).get(year))
                c = ws2.cell(row=row, column=j, value=val)
                c.font, c.border, c.alignment = _XLSX_CELL_FONT, _XLSX_BORDER, _XLSX_CENTER
            ws2.cell(row=row, column=1).border = _XLSX_BORDER
            ws2.cell(row=row, column=2).border = _XLSX_BORDER
            row += 1
    ws2.freeze_panes = "C2"
    _xlsx_autofit(ws2)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Collecte les projections économiques du RPM de la Banque du Canada.")
    parser.add_argument("--out", default="boc_mpr.json", help="Chemin du fichier JSON de sortie.")
    parser.add_argument("--csv", default=None, help="Chemin optionnel d'export CSV à plat.")
    parser.add_argument("--xlsx", default=None,
                         help="Chemin de l'export Excel (.xlsx). Par défaut : même nom que --out, en .xlsx "
                              "(l'Excel est toujours généré, sauf avec --no-xlsx).")
    parser.add_argument("--no-xlsx", action="store_true", help="Ne pas générer le fichier Excel.")
    parser.add_argument("--since", type=int, default=None, help="Année minimale à inclure (ex: 2024).")
    parser.add_argument("--latest", action="store_true", help="Ne récupérer que le RPM le plus récent.")
    parser.add_argument("--cache-dir", default=None, help="Dossier de cache des pages HTML téléchargées.")
    parser.add_argument("--pause", type=float, default=1.0, help="Pause en secondes entre deux requêtes.")
    args = parser.parse_args()

    cache_dir = Path(args.cache_dir) if args.cache_dir else None

    print("Recherche des rapports RPM disponibles...")
    dates = discover_report_dates(cache_dir=cache_dir)

    if args.since:
        dates = [d for d in dates if int(d[:4]) >= args.since]
    if args.latest and dates:
        dates = [dates[-1]]

    if not dates:
        print("Aucun RPM trouvé avec ces critères.", file=sys.stderr)
        sys.exit(1)

    print(f"{len(dates)} RPM à traiter : {', '.join(dates)}")

    reports: list[MprReport] = []
    for date in dates:
        print(f"-> {date} ...")
        report = scrape_report(date, cache_dir=cache_dir, pause=args.pause)
        if report:
            reports.append(report)

    out_path = Path(args.out)
    write_json(reports, out_path)
    print(f"JSON écrit : {out_path} ({len(reports)} rapports)")

    if args.csv:
        write_csv(reports, Path(args.csv))
        print(f"CSV écrit : {args.csv}")

    if args.xlsx or not args.no_xlsx:
        xlsx_path = Path(args.xlsx) if args.xlsx else out_path.with_suffix(".xlsx")
        write_xlsx(reports, xlsx_path)
        print(f"Excel écrit : {xlsx_path}")


if __name__ == "__main__":
    main()
