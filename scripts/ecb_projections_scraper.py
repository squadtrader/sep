#!/usr/bin/env python3
"""
ecb_projections_scraper.py
----------------------------
Collecte les projections macroéconomiques de la BCE / de l'Eurosystème,
publiées 4 fois par an (mars et septembre par le staff BCE, juin et décembre
par le staff Eurosystème).

Contrairement au SEP de la Fed ou au RPM de la BOC, chaque rapport de la BCE
affiche, pour chaque variable et chaque année :
  - la valeur du round de projections en cours (ex: "March 2026")
  - la RÉVISION (en points de pourcentage) par rapport au round précédent
    (ex: "Revisions vs December 2025"), et non la valeur précédente elle-même.

Deux tableaux sont extraits de la page de chaque rapport :
  - Table 2 : "Real GDP, trade and labour market projections" (PIB, consommation,
    investissement, exports/imports, chômage...)
  - Table 3 : "Price and cost developments for the euro area" (HICP, HICP
    sous-jacent, coûts, salaires...)

Note : la BCE publie aussi, pour chaque rapport, un fichier Excel officiel
("Projections charts and tables", lien "Annexes" sur la page de liste). Ce
fichier n'a pas pu être inspecté ici (contenu binaire, domaine non accessible
depuis cet environnement), donc ce script se base sur les tableaux HTML de la
page du rapport, dont la structure a été vérifiée directement.

Étapes :
1. Parcourt la page listant tous les rapports (all-releases.en.html) pour
   trouver l'URL de chaque round de projections.
2. Pour chaque rapport, télécharge sa page et en extrait les Tableaux 2 et 3
   (repérés par leur CONTENU — libellés de lignes attendus — plutôt que par
   leur numéro, qui peut varier selon les éditions).
3. Sauvegarde le tout en JSON, avec exports CSV et Excel automatiques.

Usage :
    python ecb_projections_scraper.py                # tous les rounds depuis décembre 2023
    python ecb_projections_scraper.py --since 2025    # seulement à partir de 2025
    python ecb_projections_scraper.py --latest
    python ecb_projections_scraper.py --out ecb_proj.json --no-xlsx
    python ecb_projections_scraper.py --cache-dir .cache_ecb

Note : les rapports antérieurs à décembre 2023 utilisent une mise en page différente
(Tableaux PIB/emploi et prix/coûts non détectés par ce parseur) et sont donc exclus
par défaut -- voir MIN_SUPPORTED_PERIOD.

Dépendances : requests, beautifulsoup4, openpyxl
    pip install requests beautifulsoup4 openpyxl
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ALL_RELEASES_URL = "https://www.ecb.europa.eu/press/projections/html/all-releases.en.html"

# Les rapports antérieurs à décembre 2023 utilisent une mise en page différente,
# où les Tableaux PIB/emploi et prix/coûts ne sont pas détectés par le parseur
# actuel (testé et confirmé : tout ce qui précède 2023-12 échoue, tout ce qui
# suit fonctionne). On se limite donc par défaut à cette période fiable.
MIN_SUPPORTED_PERIOD = "2023-12"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    )
}

# Lien vers un rapport dans le HTML de la page de liste : les hrefs y sont en
# chemin RELATIF (ex: href="/press/projections/html/ecb.projections202609_ecbstaff~8e340fc69d.en.html"),
# pas en URL absolue -> on capture le chemin puis on reconstruit l'URL complète.
REPORT_LINK_RE = re.compile(
    r'href="(/press/projections/html/'
    r'(ecb\.projections(\d{6})_(ecbstaff|eurosystemstaff)(?:~[a-f0-9]+)?)\.en\.html)"'
)
ECB_BASE = "https://www.ecb.europa.eu"

TABLE2_ROW_PATTERNS = [
    ("real_gdp", re.compile(r"^real gdp$", re.I)),
    ("private_consumption", re.compile(r"^private consumption$", re.I)),
    ("government_consumption", re.compile(r"^government consumption$", re.I)),
    ("investment", re.compile(r"^investment$", re.I)),
    ("exports", re.compile(r"^exports", re.I)),
    ("imports", re.compile(r"^imports", re.I)),
    ("domestic_demand", re.compile(r"^domestic demand$", re.I)),
    ("net_exports", re.compile(r"^net exports$", re.I)),
    ("employment", re.compile(r"^employment", re.I)),
    ("unemployment_rate", re.compile(r"^unemployment rate$", re.I)),
]

TABLE3_ROW_PATTERNS = [
    ("hicp", re.compile(r"^hicp$", re.I)),
    ("hicpx", re.compile(r"^hicp excluding energy and food$", re.I)),
    ("hicp_excl_energy", re.compile(r"^hicp excluding energy$", re.I)),
    ("hicp_energy", re.compile(r"^hicp energy$", re.I)),
    ("hicp_food", re.compile(r"^hicp food$", re.I)),
    ("gdp_deflator", re.compile(r"^gdp deflator$", re.I)),
    ("compensation_per_employee", re.compile(r"^compensation per employee$", re.I)),
    ("unit_labour_costs", re.compile(r"^unit labour costs$", re.I)),
]


@dataclass
class EcbReport:
    date: str          # AAAA-MM (mois du round de projections)
    staff_type: str     # "ecbstaff" (mars/sept) ou "eurosystemstaff" (juin/déc)
    url: str
    table2: dict = field(default_factory=dict)  # variable -> {"label", "current": {year: val}, "revision": {year: val}}
    table3: dict = field(default_factory=dict)


def fetch(url: str, cache_dir: Optional[Path] = None) -> str:
    if cache_dir:
        cache_dir.mkdir(parents=True, exist_ok=True)
        cache_file = cache_dir / (re.sub(r"[^a-zA-Z0-9]+", "_", url) + ".html")
        if cache_file.exists():
            return cache_file.read_text(encoding="utf-8")

    resp = requests.get(url, headers=HEADERS, timeout=30)
    resp.raise_for_status()
    resp.encoding = "utf-8"  # forcer UTF-8 (le site ne déclare pas toujours son charset)
    html = resp.text

    if cache_dir:
        cache_file.write_text(html, encoding="utf-8")

    return html


def discover_reports(cache_dir: Optional[Path] = None, min_period: str = MIN_SUPPORTED_PERIOD) -> list[tuple[str, str, str]]:
    """Retourne une liste de (date AAAA-MM, staff_type, url) pour chaque round trouvé,
    à partir de min_period (les rapports plus anciens ont une mise en page non prise
    en charge par ce parseur -- voir MIN_SUPPORTED_PERIOD)."""
    html = fetch(ALL_RELEASES_URL, cache_dir=cache_dir)
    seen = {}
    for m in REPORT_LINK_RE.finditer(html):
        path, yyyymm, staff_type = m.group(1), m.group(3), m.group(4)
        url = ECB_BASE + path
        date = f"{yyyymm[:4]}-{yyyymm[4:]}"
        if min_period and date < min_period:
            continue
        seen[date] = (date, staff_type, url)  # dédoublonne (même round référencé plusieurs fois : langues, etc.)

    if not seen:
        # Rien trouvé : on sauvegarde le HTML brut reçu pour pouvoir diagnostiquer
        # (blocage anti-bot, structure différente, page rendue en JavaScript, etc.)
        debug_path = Path("debug_ecb_all_releases.html")
        debug_path.write_text(html, encoding="utf-8")
        print(f"  ! Aucun lien de rapport trouvé dans la page. HTML brut sauvegardé dans "
              f"{debug_path.resolve()} ({len(html)} caractères) pour diagnostic.", file=sys.stderr)
        snippet = re.sub(r"\s+", " ", html)[:500]
        print(f"  ! Début du contenu reçu : {snippet}", file=sys.stderr)

    return sorted(seen.values())


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _find_table_by_content(soup: BeautifulSoup, patterns: list, min_matches: int):
    """Repère la <table> voulue par son contenu (libellés de lignes attendus),
    car le numéro de tableau ("Table 2", "Table 3") peut varier selon les
    éditions (nombre de graphiques/tableaux différent en amont sur la page)."""
    best_table, best_score = None, 0
    for table in soup.find_all("table"):
        matched_keys = set()
        for row in table.find_all("tr"):
            cells = row.find_all(["th", "td"])
            if not cells:
                continue
            label = _clean(cells[0].get_text(" "))
            if not label:
                continue
            for key, pattern in patterns:
                if pattern.search(label):
                    matched_keys.add(key)
        if len(matched_keys) > best_score:
            best_table, best_score = table, len(matched_keys)
    return best_table if best_score >= min_matches else None


def _expand_row(row) -> list[str]:
    out = []
    for cell in row.find_all(["th", "td"]):
        text = _clean(cell.get_text(" "))
        colspan = int(cell.get("colspan", 1))
        out.extend([text] * colspan)
    return out


def _parse_table(table, patterns: list) -> dict:
    """Parse une table à deux lignes d'en-tête : bloc 'valeurs actuelles' (années)
    puis bloc 'révisions vs round précédent' (mêmes années), 4 colonnes chacun."""
    rows = table.find_all("tr")
    if len(rows) < 3:
        return {}

    year_row = _expand_row(rows[1])  # ligne 2 = années (répétées pour les 2 blocs)
    years = [y for y in year_row if re.fullmatch(r"(19|20)\d{2}", y)]
    n_years = len(years) // 2 if len(years) % 2 == 0 and len(years) >= 2 else len(years)
    current_years = years[:n_years]
    revision_years = years[n_years:2 * n_years] or current_years

    result: dict = {}
    for row in rows[2:]:
        cells = row.find_all(["th", "td"])
        if not cells:
            continue
        texts = [_clean(c.get_text(" ")) for c in cells]
        if not texts or not texts[0]:
            continue
        label = texts[0]

        matched_key = None
        for key, pattern in patterns:
            if pattern.search(label):
                matched_key = key
                break
        if not matched_key:
            continue

        data_cells = texts[1:]
        current = {}
        for i, year in enumerate(current_years):
            if i < len(data_cells) and data_cells[i]:
                current[year] = data_cells[i]
        revision = {}
        offset = n_years
        for i, year in enumerate(revision_years):
            idx = offset + i
            if idx < len(data_cells) and data_cells[idx]:
                revision[year] = data_cells[idx]

        result[matched_key] = {"label": label, "current": current, "revision": revision}

    return result


def scrape_report(date: str, staff_type: str, url: str, cache_dir: Optional[Path] = None,
                   pause: float = 1.0) -> Optional[EcbReport]:
    try:
        html = fetch(url, cache_dir=cache_dir)
    except requests.HTTPError as exc:
        print(f"  ! {date}: page introuvable ({exc})", file=sys.stderr)
        return None

    soup = BeautifulSoup(html, "html.parser")
    report = EcbReport(date=date, staff_type=staff_type, url=url)

    table2 = _find_table_by_content(soup, TABLE2_ROW_PATTERNS, min_matches=4)
    if table2 is not None:
        report.table2 = _parse_table(table2, TABLE2_ROW_PATTERNS)
    else:
        print(f"  ! {date}: tableau PIB/emploi introuvable", file=sys.stderr)

    table3 = _find_table_by_content(soup, TABLE3_ROW_PATTERNS, min_matches=3)
    if table3 is not None:
        report.table3 = _parse_table(table3, TABLE3_ROW_PATTERNS)
    else:
        print(f"  ! {date}: tableau prix/coûts introuvable", file=sys.stderr)

    if table2 is None or table3 is None:
        debug_path = Path(f"debug_ecb_report_{date}.html")
        debug_path.write_text(html, encoding="utf-8")
        print(f"    -> HTML brut sauvegardé dans {debug_path.resolve()} pour diagnostic.", file=sys.stderr)

    if not report.table2 and not report.table3:
        return None

    time.sleep(pause)
    return report


def report_to_dict(r: EcbReport) -> dict:
    return {
        "date": r.date,
        "staff_type": r.staff_type,
        "url": r.url,
        "table2_gdp_labour": r.table2,
        "table3_prices_costs": r.table3,
    }


def write_json(reports: list[EcbReport], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    payload = [report_to_dict(r) for r in reports]
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(reports: list[EcbReport], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["report_date", "table", "variable", "label", "year", "current_value", "revision"])
        for r in reports:
            for table_name, table_data in (("table2", r.table2), ("table3", r.table3)):
                for var_key, var_data in table_data.items():
                    label = var_data.get("label", var_key)
                    years = set(var_data.get("current", {})) | set(var_data.get("revision", {}))
                    for year in sorted(years):
                        writer.writerow([r.date, table_name, var_key, label, year,
                                          var_data.get("current", {}).get(year, ""),
                                          var_data.get("revision", {}).get(year, "")])


# --- Export Excel -----------------------------------------------------------

_XLSX_FONT_NAME = "Arial"
_XLSX_HEADER_FILL = PatternFill("solid", fgColor="003399")  # bleu BCE
_XLSX_HEADER_FONT = Font(name=_XLSX_FONT_NAME, size=10, bold=True, color="FFFFFF")
_XLSX_SUBHEADER_FILL = PatternFill("solid", fgColor="D6E0F5")
_XLSX_SUBHEADER_FONT = Font(name=_XLSX_FONT_NAME, size=10, bold=True)
_XLSX_CELL_FONT = Font(name=_XLSX_FONT_NAME, size=10)
_XLSX_THIN = Side(style="thin", color="B7B7B7")
_XLSX_BORDER = Border(left=_XLSX_THIN, right=_XLSX_THIN, top=_XLSX_THIN, bottom=_XLSX_THIN)
_XLSX_CENTER = Alignment(horizontal="center", vertical="center")

SUMMARY_VARIABLES = [
    ("real_gdp", "PIB réel"),
    ("hicp", "Inflation IPCH"),
    ("hicpx", "Inflation IPCH sous-jacente"),
    ("unemployment_rate", "Taux de chômage"),
]
_SUMMARY_TABLE_OF = {"real_gdp": "table2", "unemployment_rate": "table2", "hicp": "table3", "hicpx": "table3"}


def _xlsx_style_header(cell, fill, font) -> None:
    cell.fill, cell.font, cell.alignment, cell.border = fill, font, _XLSX_CENTER, _XLSX_BORDER


def _xlsx_autofit(ws, min_width=10, max_width=24) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, min_width), max_width)


def _fmt_cell(var_data: dict, year: str) -> str:
    current = var_data.get("current", {}).get(year)
    revision = var_data.get("revision", {}).get(year)
    if current is None:
        return ""
    if revision in (None, ""):
        return current
    # La révision est déjà un texte signé tel qu'affiché sur le site (ex: "-0.3", "0.1")
    sign_prefix = "" if revision.startswith(("-", "+")) else "+"
    return f"{current} ({sign_prefix}{revision} pp)"


def write_xlsx(reports: list[EcbReport], out_path: Path) -> None:
    if not reports:
        return
    reports = sorted(reports, key=lambda r: r.date)

    years: set = set()
    for r in reports:
        for table_data in (r.table2, r.table3):
            for var_data in table_data.values():
                years.update(var_data.get("current", {}).keys())
    years_sorted = sorted(years, key=lambda y: int(y))

    wb = Workbook()
    ws = wb.active
    ws.title = "Résumé"
    ws.cell(row=1, column=1, value="Rapport")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    _xlsx_style_header(ws.cell(row=1, column=1), _XLSX_HEADER_FILL, _XLSX_HEADER_FONT)

    col = 2
    for var_key, var_label in SUMMARY_VARIABLES:
        start_col = col
        for year in years_sorted:
            c = ws.cell(row=2, column=col, value=year)
            _xlsx_style_header(c, _XLSX_SUBHEADER_FILL, _XLSX_SUBHEADER_FONT)
            col += 1
        if col == start_col:
            c = ws.cell(row=2, column=col, value="N/D")
            _xlsx_style_header(c, _XLSX_SUBHEADER_FILL, _XLSX_SUBHEADER_FONT)
            col += 1
        if col - 1 > start_col:
            ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=col - 1)
        _xlsx_style_header(ws.cell(row=1, column=start_col, value=var_label), _XLSX_HEADER_FILL, _XLSX_HEADER_FONT)

    row = 3
    for r in reports:
        label = f"{r.date} ({'Eurosystème' if r.staff_type == 'eurosystemstaff' else 'BCE'})"
        cell = ws.cell(row=row, column=1, value=label)
        cell.font, cell.border, cell.alignment = _XLSX_CELL_FONT, _XLSX_BORDER, _XLSX_CENTER
        col = 2
        for var_key, _ in SUMMARY_VARIABLES:
            table_data = r.table2 if _SUMMARY_TABLE_OF[var_key] == "table2" else r.table3
            var_data = table_data.get(var_key, {})
            for year in years_sorted:
                c = ws.cell(row=row, column=col, value=_fmt_cell(var_data, year) if var_data else "")
                c.font, c.border, c.alignment = _XLSX_CELL_FONT, _XLSX_BORDER, _XLSX_CENTER
                col += 1
        row += 1

    ws.freeze_panes = "B3"
    _xlsx_autofit(ws)
    note_row = row + 1
    ws.cell(row=note_row, column=1,
            value="Valeur du round en cours ; entre parenthèses, la révision en points de pourcentage "
                  "par rapport au round précédent (et non la valeur précédente elle-même).")
    ws.cell(row=note_row, column=1).font = Font(name=_XLSX_FONT_NAME, size=9, italic=True)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Collecte les projections macroéconomiques de la BCE.")
    parser.add_argument("--out", default="ecb_projections.json", help="Chemin du fichier JSON de sortie.")
    parser.add_argument("--csv", default=None, help="Chemin optionnel d'export CSV à plat.")
    parser.add_argument("--xlsx", default=None,
                         help="Chemin de l'export Excel (.xlsx). Par défaut : même nom que --out, en .xlsx.")
    parser.add_argument("--no-xlsx", action="store_true", help="Ne pas générer le fichier Excel.")
    parser.add_argument("--since", type=int, default=None, help="Année minimale à inclure (ex: 2023).")
    parser.add_argument("--latest", action="store_true", help="Ne récupérer que le round le plus récent.")
    parser.add_argument("--cache-dir", default=None, help="Dossier de cache des pages HTML téléchargées.")
    parser.add_argument("--pause", type=float, default=1.0, help="Pause en secondes entre deux requêtes.")
    args = parser.parse_args()

    cache_dir = Path(args.cache_dir) if args.cache_dir else None

    print("Recherche des rounds de projections BCE disponibles...")
    entries = discover_reports(cache_dir=cache_dir)

    if args.since:
        entries = [e for e in entries if int(e[0][:4]) >= args.since]
    if args.latest and entries:
        entries = [entries[-1]]

    if not entries:
        print("Aucun round de projections trouvé avec ces critères.", file=sys.stderr)
        sys.exit(1)

    print(f"{len(entries)} round(s) à traiter : {', '.join(e[0] for e in entries)}")

    reports: list[EcbReport] = []
    for date, staff_type, url in entries:
        print(f"-> {date} ({staff_type}) ...")
        report = scrape_report(date, staff_type, url, cache_dir=cache_dir, pause=args.pause)
        if report:
            reports.append(report)

    out_path = Path(args.out)
    write_json(reports, out_path)
    print(f"JSON écrit : {out_path} ({len(reports)} rounds)")

    if args.csv:
        write_csv(reports, Path(args.csv))
        print(f"CSV écrit : {args.csv}")

    if not args.no_xlsx:
        xlsx_path = Path(args.xlsx) if args.xlsx else out_path.with_suffix(".xlsx")
        write_xlsx(reports, xlsx_path)
        print(f"Excel écrit : {xlsx_path}")


if __name__ == "__main__":
    main()
